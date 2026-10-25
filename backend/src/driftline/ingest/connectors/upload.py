"""Bring-your-own documents: a local folder of Markdown, PDF and DOCX files.

The external id is the file's path relative to the root, and the version key is the
sha256 of its bytes, so re-uploading a changed file at the same path is a new version.
"""

import hashlib
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from driftline.ingest.markdown import parse_markdown
from driftline.ingest.types import (
    RawDocument,
    Section,
    SourceFetchError,
    render_content,
    unique_sections,
)

SUPPORTED_SUFFIXES = (".md", ".markdown", ".pdf", ".docx")


def parse_pdf(data: bytes) -> tuple[str | None, tuple[Section, ...]]:
    """One section per page (`page 1`, `page 2`, ...); title from the PDF metadata."""
    try:
        reader = PdfReader(BytesIO(data))
        pages = [
            (f"page {i}", (page.extract_text() or "").strip())
            for i, page in enumerate(reader.pages, start=1)
        ]
        title = reader.metadata.title if reader.metadata and reader.metadata.title else None
    except PdfReadError as exc:
        raise SourceFetchError(f"unreadable PDF: {exc}", retryable=False) from exc
    return title, unique_sections(pages)


def parse_docx(data: bytes) -> tuple[str | None, tuple[Section, ...]]:
    """Sections follow the `Heading N` paragraph styles, like Markdown headings."""
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated exception types
        raise SourceFetchError(f"unreadable DOCX: {exc}", retryable=False) from exc
    stack: list[tuple[int, str]] = []
    pairs: list[tuple[str, str]] = []
    path, buf = "", []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if style.startswith("Heading ") and style[8:].isdigit() and text:
            pairs.append((path, "\n".join(buf).strip()))
            level = int(style[8:])
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, text))
            path, buf = " > ".join(t for _, t in stack), []
        elif text:
            buf.append(text)
    pairs.append((path, "\n".join(buf).strip()))
    title = document.core_properties.title or None
    return title, unique_sections([(p, t) for p, t in pairs if p or t])


class UploadConnector:
    kind = "upload"

    def __init__(self, root: Path | str) -> None:
        self._root = Path(root)

    async def list_external_ids(self) -> list[str]:
        if not self._root.is_dir():
            raise SourceFetchError(f"upload root {self._root} is not a directory", retryable=False)
        return sorted(
            p.relative_to(self._root).as_posix()
            for p in self._root.rglob("*")
            if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
        )

    def _resolve(self, external_id: str) -> Path:
        path = (self._root / external_id).resolve()
        if not path.is_relative_to(self._root.resolve()):
            raise SourceFetchError(f"{external_id!r} escapes the upload root", retryable=False)
        if not path.is_file():
            raise SourceFetchError(f"{external_id!r} not found under {self._root}", retryable=False)
        return path

    async def fetch(
        self, external_id: str, *, since: str | None = None
    ) -> AsyncIterator[RawDocument]:
        """The file's current content, unless its sha256 equals the `since` cursor."""
        path = self._resolve(external_id)
        data = path.read_bytes()
        version_key = hashlib.sha256(data).hexdigest()
        if version_key == since:
            return
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            title, sections = parse_pdf(data)
        elif suffix == ".docx":
            title, sections = parse_docx(data)
        else:
            doc = parse_markdown(data.decode("utf-8", errors="replace"))
            title, sections = doc.title, doc.sections
        stat = path.stat()
        yield RawDocument(
            external_id=external_id,
            title=title or path.stem,
            version_key=version_key,
            valid_from=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
            content=render_content(sections),
            sections=sections,
            raw_uri=path.as_uri(),
            source_meta={"size": stat.st_size, "suffix": suffix},
        )
