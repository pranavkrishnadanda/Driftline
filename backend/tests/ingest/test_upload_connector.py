import io
import os
from datetime import UTC, datetime

import docx
import pytest

from driftline.ingest.connectors.upload import UploadConnector
from driftline.ingest.types import SourceFetchError


def minimal_pdf(pages: list[str]) -> bytes:
    """A valid PDF with one line of Helvetica text per page, built without a PDF writer."""
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET"
        objs.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


async def fetch_all(connector: UploadConnector, external_id: str, since: str | None = None):
    return [d async for d in connector.fetch(external_id, since=since)]


async def test_lists_supported_files_sorted(tmp_path):
    (tmp_path / "guides").mkdir()
    for name in ["b.md", "guides/a.markdown", "c.pdf", "d.docx", "notes.txt"]:
        (tmp_path / name).write_bytes(b"x")
    assert await UploadConnector(tmp_path).list_external_ids() == [
        "b.md",
        "c.pdf",
        "d.docx",
        "guides/a.markdown",
    ]


async def test_markdown_file_becomes_one_version(tmp_path):
    path = tmp_path / "runbook.md"
    path.write_text("# Rollout runbook\n\n## Undo\n\nkubectl rollout undo.\n")
    os.utime(path, (1_700_000_000, 1_700_000_000))
    [doc] = await fetch_all(UploadConnector(tmp_path), "runbook.md")
    assert doc.title == "Rollout runbook"
    assert [s.path for s in doc.sections] == ["Rollout runbook", "Rollout runbook > Undo"]
    assert doc.valid_from == datetime.fromtimestamp(1_700_000_000, tz=UTC)
    assert len(doc.version_key) == 64
    assert doc.raw_uri == path.resolve().as_uri()
    assert doc.external_id == "runbook.md"


async def test_unchanged_file_with_cursor_yields_nothing(tmp_path):
    (tmp_path / "a.md").write_text("# A\n\nbody\n")
    connector = UploadConnector(tmp_path)
    [doc] = await fetch_all(connector, "a.md")
    assert await fetch_all(connector, "a.md", since=doc.version_key) == []


async def test_changed_file_gets_new_version_key(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# A\n\nbody\n")
    connector = UploadConnector(tmp_path)
    [first] = await fetch_all(connector, "a.md")
    path.write_text("# A\n\nbody changed\n")
    [second] = await fetch_all(connector, "a.md", since=first.version_key)
    assert second.version_key != first.version_key
    assert second.sections[0].text == "body changed"


async def test_pdf_pages_become_sections(tmp_path):
    pdf = minimal_pdf(["Boxed warning: liver (hepatic) failure.", "Page two text."])
    (tmp_path / "label.pdf").write_bytes(pdf)
    [doc] = await fetch_all(UploadConnector(tmp_path), "label.pdf")
    assert [s.path for s in doc.sections] == ["page 1", "page 2"]
    assert doc.sections[0].text == "Boxed warning: liver (hepatic) failure."
    assert doc.title == "label"


async def test_docx_headings_become_section_paths(tmp_path):
    document = docx.Document()
    document.core_properties.title = "Runbook"
    document.add_paragraph("Preamble line.")
    document.add_heading("Rollout", level=1)
    document.add_paragraph("Use kubectl rollout.")
    document.add_heading("Undo", level=2)
    document.add_paragraph("kubectl rollout undo.")
    document.save(tmp_path / "runbook.docx")
    [doc] = await fetch_all(UploadConnector(tmp_path), "runbook.docx")
    assert doc.title == "Runbook"
    assert [(s.path, s.text) for s in doc.sections] == [
        ("", "Preamble line."),
        ("Rollout", "Use kubectl rollout."),
        ("Rollout > Undo", "kubectl rollout undo."),
    ]


async def test_path_escape_is_rejected(tmp_path):
    (tmp_path / "root").mkdir()
    (tmp_path / "secret.md").write_text("# no\n")
    with pytest.raises(SourceFetchError, match="escapes"):
        await fetch_all(UploadConnector(tmp_path / "root"), "../secret.md")


async def test_missing_root_and_file_raise_source_fetch_error(tmp_path):
    with pytest.raises(SourceFetchError):
        await UploadConnector(tmp_path / "nope").list_external_ids()
    with pytest.raises(SourceFetchError) as info:
        await fetch_all(UploadConnector(tmp_path), "missing.md")
    assert info.value.retryable is False
